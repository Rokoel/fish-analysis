from dataclasses import asdict, replace
import json
from types import SimpleNamespace

import numpy as np
import pytest
from scipy import ndimage as ndi

from fish_viewer.refinement.config import RefinementConfig, read_model
from fish_viewer.refinement.data import sample_target
from fish_viewer.refinement.inference import resample_labels, refine, surface_evidence


def test_finer_labels_preserve_physical_centers_and_instance_ids():
    coarse=np.zeros((8,)*3,np.uint32);coarse[2:6,2:6,2:6]=7
    fine=resample_labels(coarse,np.array([.3]*3),np.array([.15]*3),
                         (16,)*3,(.15,)*3,(.075,)*3)
    assert set(np.unique(fine))=={0,7}
    assert np.array_equal(np.argwhere(fine==7).mean(0)*.15+.075,[1.2]*3)
    assert np.count_nonzero(fine==7)==8*np.count_nonzero(coarse==7)


def test_target_crop_labels_neighbor_as_other_instance_and_ignores_padding():
    config=replace(RefinementConfig(),patch_size=32,margin_um=.8)
    image=np.zeros((24,)*3,np.float32)
    box=(slice(1,9),slice(1,9),slice(1,9));mask=np.ones((8,)*3,bool)
    image[box]=.2;image[12:20,1:9,1:9]=1.  # unannotated neighbor, not global negative
    sample,truth,weight,valid,seed,step=sample_target(image,box,mask,(.2,)*3,config)
    assert not truth[sample>.8].any()  # not a member of this chosen target
    assert truth.any() and (~valid).any()
    assert (weight[~valid]==0).all()
    assert truth[tuple(np.rint(seed).astype(int))]
    assert step>0


def test_surface_normals_follow_visible_concave_boundary_and_reject_blank():
    grid=np.indices((48,)*3);radius=np.linalg.norm((grid-23.5)*.2,axis=0)
    mask=radius<3.
    mask[16:32,16:32,30:]=False  # a visible indentation
    boundary=mask & ~ndi.binary_erosion(mask)
    image=np.exp(-(ndi.distance_transform_edt(~boundary,sampling=.2)/.2)**2).astype(np.float32)
    visible=surface_evidence(mask,image,(.2,)*3)
    blank=surface_evidence(mask,np.zeros_like(image),(.2,)*3)
    assert visible['surface_support']>.7
    assert blank['surface_support']==0


def test_refinement_can_indent_a_cell_without_adding_instances():
    config=replace(RefinementConfig(),patch_size=32,margin_um=1.)
    shape=(40,)*3;grid=np.indices(shape);radius=np.linalg.norm(grid-19.5,axis=0)
    labels=(radius<8).astype(np.uint32)
    image=np.exp(-((radius-8)/1.2)**2).astype(np.float32)
    prepared=SimpleNamespace(image=image,spacing=(.2,)*3)
    class IndentedModel:
        def __call__(self,x,training=False):
            g=np.indices((32,)*3);r=np.linalg.norm(g-15.5,axis=0)
            target=r<8.7
            target[13:19,13:19,21:]=False
            return np.where(target,12.,-12.)[None,...,None]
    result,diagnostics=refine(labels,prepared,IndentedModel(),config,evidence_threshold=0.)
    assert set(np.unique(result))=={0,1}
    assert diagnostics[0]['accepted']
    assert np.count_nonzero(result)<np.count_nonzero(labels)
    class EmptyModel:
        def __call__(self,x,training=False):return np.full((1,32,32,32,1),-20.)
    fallback,diagnostics=refine(labels,prepared,EmptyModel(),config,evidence_threshold=0.)
    assert np.array_equal(fallback,labels)
    assert not diagnostics[0]['accepted'] and 'seed' in diagnostics[0]['fallback_reason']


def test_refinement_manifest_rejects_invalid_checkpoint_and_config(tmp_path):
    manifest=dict(version=1,engine='seeded-unet3d',config=asdict(RefinementConfig()),checkpoint='best.weights.h5')
    (tmp_path/'refinement-model.json').write_text(json.dumps(manifest))
    with pytest.raises(ValueError,match='checkpoint'):read_model(tmp_path)
    (tmp_path/'best.weights.h5').write_bytes(b'placeholder')
    assert read_model(tmp_path)['engine']=='seeded-unet3d'
    with pytest.raises(ValueError,match='finite'):replace(RefinementConfig(),voxel_um=float('nan')).validate()


def test_finer_cache_application_preserves_saved_geometry_without_tensorflow(tmp_path,monkeypatch):
    import fish_viewer.refinement.inference as worker
    shape=(24,)*3;labels=np.zeros(shape,np.uint32);labels[8:16,8:16,8:16]=1
    image=np.zeros(shape,np.float32);np.save(tmp_path/'image.npy',image)
    config=asdict(RefinementConfig());model=tmp_path/'model';model.mkdir()
    (model/'refinement-model.json').write_text(json.dumps(dict(version=1,engine='seeded-unet3d',
        checkpoint='best.weights.h5',config=config)))
    (model/'best.weights.h5').write_bytes(b'placeholder')
    metadata=dict(version=1,image='scan.ims',time=0,channel=0,channel_scores=[],fingerprint='same',
                  spacing_zyx=[.15]*3,origin_xyz=[.075]*3,config=dict(rays=64,use_shape_filter=False))
    prediction=tmp_path/'cells.npz';np.savez_compressed(prediction,labels=labels,metadata=np.array(json.dumps(metadata)))
    source=tmp_path/'cells.json';source.write_text(json.dumps(dict(status='complete',files=[dict(image='scan.ims',predictions=str(prediction))])))
    audit=tmp_path/'audit.json';audit.write_text(json.dumps(dict(status='complete',cases=[dict(image='scan.ims',directory=str(tmp_path),
        audit=dict(fingerprint='same',channel=0,spacing_zyx=[.15]*3,origin_xyz=[.075]*3))])))
    monkeypatch.setattr(worker,'runtime',lambda *a:(None,None,None,{}))
    monkeypatch.setattr(worker,'network',lambda *a:SimpleNamespace(load_weights=lambda p:None))
    monkeypatch.setattr(worker,'refine',lambda labels,*a:(labels,[dict(cell_id=1,accepted=False)]))
    report=worker.apply(source,model,tmp_path/'output',audit=audit)
    assert report['status']=='complete',report.get('error')
    with np.load(report['files'][0]['predictions'],allow_pickle=False) as archive:
        saved=json.loads(str(archive['metadata']))
        assert np.array_equal(archive['labels'],labels)
        assert saved['spacing_zyx']==[.15]*3 and saved['refinement']['cells'][0]['cell_id']==1


def test_excel_remains_two_lists_with_refinement_metadata(tmp_path):
    from openpyxl import load_workbook
    from fish_viewer.workflow.excel import write_excel
    report=dict(points=[],counts=[],cell_volumes=[dict(cell_id=1,volume_um3=10.)],
                boundary_refinement=dict(model='example',cells=[dict(cell_id=1,accepted=False,fallback_reason='Weak membrane')]))
    path=write_excel(tmp_path/'report.xlsx',report)
    book=load_workbook(path,read_only=True,data_only=True)
    assert book.sheetnames == ['Measurements']
    assert list(book['Measurements'].values)[1] == ('Program',1,10.,None,None,None,None)
    book.close()


@pytest.mark.gui
def test_boundary_training_live_epoch_and_graph_progress(qapp,tmp_path):
    from fish_viewer.refinement.gui import TrainingDialog
    dialog=TrainingDialog(output=str(tmp_path))
    try:
        dialog.handle_event(dict(kind='progress',phase='Training boundary network',epoch=12,epochs=60,
                                 current=42,total=859,loss=.15,dice=.91))
        assert '12 / 60' in dialog.epoch.text()
        assert dialog.progress.value()==42 and dialog.progress.maximum()==859
        dialog.handle_event(dict(kind='epoch',epoch=12,training=dict(loss=.15,dice=.91),
                                 validation=dict(loss=.19,dice=.89)))
        assert len(dialog.rows)==1 and len(dialog.axes[0].lines)==2
    finally:dialog.close();qapp.processEvents()


@pytest.mark.gui
def test_analysis_refinement_progress_has_six_stages(qapp,tmp_path):
    from PySide6 import QtCore
    from fish_viewer.analysis.app import AnalysisWindow
    settings=QtCore.QSettings(str(tmp_path/'settings.ini'),QtCore.QSettings.Format.IniFormat)
    window=AnalysisWindow(off_screen=True,settings=settings)
    try:
        window.scan_number=1;window.total_scans=1;window.active_image='example.ims'
        window.processing_progress(dict(stage=2,stages=6,stage_name='Refining cell boundaries',
                                         phase='Refining cell boundaries',current=5,total=10))
        assert '2 / 6' in window.phase.text()
        assert window.overall.value()==1250
        assert window.boundary_model.placeholderText()=='Optional 3D U-Net model folder'
    finally:window.close();qapp.processEvents()


def test_redrawn_proposals_never_fall_back_to_detector_masks():
    from fish_viewer.refinement.inference import redraw_proposals
    labels=np.zeros((32,)*3,np.uint32);labels[10:22,10:22,10:22]=7
    prepared=SimpleNamespace(image=np.zeros(labels.shape,np.float32),spacing=(.3,)*3,origin=(1.,2.,3.))
    config=replace(RefinementConfig(),patch_size=32)
    class EmptyModel:
        def __call__(self,x,training=False):return np.full((1,32,32,32,1),-20.)
    result,diagnostics=redraw_proposals(labels,prepared,EmptyModel(),config)
    assert not result.any()
    assert len(diagnostics)==1 and not diagnostics[0]['accepted']
    assert diagnostics[0]['proposal_id']==7
    assert np.allclose(diagnostics[0]['box_low_xyz'],[4.,5.,6.])
    assert np.array_equal(labels[10:22,10:22,10:22],np.full((12,)*3,7))


def test_redrawn_positive_network_output_on_blank_image_is_rejected():
    from fish_viewer.refinement.inference import redraw_proposals
    labels=np.zeros((40,)*3,np.uint32);labels[15:25,15:25,15:25]=1
    prepared=SimpleNamespace(image=np.zeros(labels.shape,np.float32),spacing=(.3,)*3,origin=(0.,)*3)
    config=replace(RefinementConfig(),patch_size=32)
    class HallucinatedCell:
        def __call__(self,x,training=False):
            r=np.linalg.norm(np.indices((32,)*3)-15.5,axis=0)
            return np.where(r<9,20.,-20.)[None,...,None]
    result,diagnostics=redraw_proposals(labels,prepared,HallucinatedCell(),config)
    assert not result.any()
    assert diagnostics[0]['rejection_reason']=='Insufficient visible membrane evidence'

from fish_viewer.refinement.training import batch,validation_boundaries
def test_surface_targets_follow_augmented_geometry_and_physical_spacing():
    size=32
    image=np.zeros((1,size,size,size),np.float32)
    mask=np.zeros_like(image);mask[:,8:24,8:24,8:24]=1
    arrays=dict(images=image,masks=mask,weights=np.ones_like(image),valid=np.ones_like(image))
    records=[dict(seed_patch=[15.5]*3,patch_spacing_um=.2)]
    config=replace(RefinementConfig(),patch_size=size,surface_loss_weight=.1,crop_scale_jitter=.1,crop_shift_um=.2)
    _,y=batch(arrays,records,[0],config,np.random.default_rng(2),augment=True)
    truth=y[0,...,0]>.5
    expected=(ndi.distance_transform_edt(~truth,sampling=.2)-ndi.distance_transform_edt(truth,sampling=.2))/2
    assert y.shape[-1]==4
    assert np.allclose(y[0,...,3],np.clip(expected,-1,1))

@pytest.mark.neural
def test_surface_loss_and_asymmetric_loss_ignore_unknown_and_penalize_missing_voxels():
    import tensorflow as tf
    from fish_viewer.refinement.network import loss
    truth=np.zeros((1,8,8,8,4),np.float32)
    truth[...,1]=1;truth[:,2:6,2:6,2:6,2]=1
    truth[:,2:6,2:6,2:6,0]=1;truth[...,3]=1;truth[...,3][truth[...,0]>0]=-1
    logits=tf.Variable(tf.zeros((1,8,8,8,1)))
    config=replace(RefinementConfig(),surface_loss_weight=.1,tversky_beta=.6)
    with tf.GradientTape() as tape:value,_=loss(tf,tf.constant(truth),logits,config)
    gradient=tape.gradient(value,logits).numpy()
    assert not gradient[:,:2].any()
    assert gradient[0,3,3,3,0]<0
    # Shift logits inside marked target: covering it should lower the objective.
    grown=np.zeros(logits.shape,np.float32);grown[:,2:6,2:6,2:6]=2
    assert float(loss(tf,truth,grown,config)[0])<float(value)

def test_validation_boundary_scores_empty_prediction_as_failure():
    truth=np.zeros((1,16,16,16,4),np.float32);truth[:,4:12,4:12,4:12,0]=1;truth[...,2]=1
    record=[dict(seed_patch=[8,8,8],patch_spacing_um=.2)]
    config=RefinementConfig()
    good=validation_boundaries(truth[...,0:1],truth,[0],record,config)[0]
    bad=validation_boundaries(np.zeros((1,16,16,16,1)),truth,[0],record,config)[0]
    assert good==(0.,0.,0.)
    assert bad[0]>1 and bad[1]==100

def test_invalid_new_loss_config_is_rejected():
    with pytest.raises(ValueError):replace(RefinementConfig(),tversky_beta=.9).validate()
    with pytest.raises(ValueError):replace(RefinementConfig(),surface_loss_weight=-.1).validate()
