import numpy as np
import pytest
from openpyxl import load_workbook

from fish_viewer.models import Surface
from fish_viewer.workflow.volumes import surface_volume
from fish_viewer.workflow.excel import write_excel


def box_surface():
    vertices = np.array([[0,0,0], [2,0,0], [2,3,0], [0,3,0],
                         [0,0,4], [2,0,4], [2,3,4], [0,3,4]], dtype=float)
    faces = np.array([[0,2,1], [0,3,2], [4,5,6], [4,6,7], [0,1,5], [0,5,4],
                      [1,2,6], [1,6,5], [2,3,7], [2,7,6], [3,0,4], [3,4,7]])
    return Surface('Predicted cell 1', vertices, faces)


def test_mesh_volume_is_physical_translation_invariant_and_rejects_open_mesh():
    surface = box_surface()
    assert surface_volume(surface) == pytest.approx(24)
    surface.vertices += [100, -70, 20]
    surface.triangles[::2] = surface.triangles[::2, ::-1]  # Inconsistent input winding is repaired.
    assert surface_volume(surface) == pytest.approx(24)
    surface.triangles = surface.triangles[:-1]
    assert surface_volume(surface) is None


def test_two_lists_keep_manual_nn_values_channels_and_safe_text(tmp_path):
    report = dict(image='scan.ims', channels=[dict(name='Membrane'),dict(name='=BAD()'),dict(name='Red')],
        points=[dict(signal_id=1,cell_id=4,channel=1,score=.999,xyz=[1,2,3],border_xyz=[1,2,4],distance_um=1),
                dict(signal_id=2,cell_id=4,channel=2,score=.998,xyz=[2,3,4],border_xyz=[2,3,6],distance_um=2)],
        counts=[dict(cell_id=4,channel=1,channel_name='=BAD()',count=1),dict(cell_id=4,channel=2,channel_name='Red',count=1)],
        cell_volumes=[dict(cell_id=4,volume_um3=24,voxel_volume_um3=25,volume_status='closed_mesh')],
        manual_measurements=dict(cells=[dict(cell_id='Manual cell',volume_um3=22)],
            points=[dict(signal_id='A1',cell_id='Manual cell',channel=None,channel_name=None,distance_um=.8)]))
    path = write_excel(tmp_path/'report.xlsx', report)
    workbook = load_workbook(path)
    assert workbook.sheetnames == ['Measurements']
    rows=list(workbook['Measurements'].values)
    assert len(rows[0]) == 7
    assert rows[1] == ('Manual','Manual cell',22,'A1',None,None,.8)
    assert rows[2] == ('Program',4,24,1,2,'=BAD()',1)
    assert rows[3] == ('Program',4,24,2,3,'Red',2)
    assert workbook['Measurements']['F3'].data_type == 's'
    workbook.close()


def test_imported_scene_is_preexisting_regardless_of_old_provenance(ims_file):
    from fish_viewer.workflow.manual import manual_measurements
    surface = box_surface()
    surface.vertices = surface.vertices * [4,2,1] + [10,20,30]
    vertices = ' '.join(str(v) for xyz in surface.vertices for v in [*xyz,0,0,1])
    faces = ' '.join(str(v) for v in surface.triangles.ravel())
    scene = ims_file.with_suffix('.imx')
    scene.write_text(f'''<bpImarisApplication><bpSurfacesViewer>
    <bpSurfaceComponent><name>Manual cell</name></bpSurfaceComponent>
    <bpSurfaces><surface><vertices>{vertices}</vertices><triangles>{faces}</triangles></surface></bpSurfaces>
    </bpSurfacesViewer>
    <bpMeasurementPoints><bpSurfaceComponent><name>Manual AB</name></bpSurfaceComponent>
    <Point Name="A" Position="13 23 31"/><Point Name="B" Position="13 23 30"/>
    </bpMeasurementPoints>
    <bpMeasurementPoints><bpSurfaceComponent><name>Manual A only</name></bpSurfaceComponent>
    <Point Name="A" Position="15 23 31"/></bpMeasurementPoints>
    <bpMeasurementPoints><Point Name="A" Position="13 23 31" TimeIndex="1"/>
    <Point Name="B" Position="13 23 30" TimeIndex="1"/></bpMeasurementPoints>
    </bpImarisApplication>''')
    result = manual_measurements(dict(image=str(ims_file),membrane_channel=0))
    assert len(result['cells']) == 1
    assert result['cells'][0]['cell_name'] == 'Manual cell'
    assert result['cells'][0]['volume_um3'] == pytest.approx(192)
    assert len(result['points']) == 2
    assert all(p['cell_id']==result['cells'][0]['cell_id'] and p['distance_um']==pytest.approx(1) for p in result['points'])
    assert all(p['channel'] is None for p in result['points'])
    # A previous export is pre-existing input to the next run, even with old provenance.
    scene.write_text(scene.read_text().replace('<bpImarisApplication>',
        '<bpImarisApplication><!--FISH Viewer exported detections; synthetic test-->'))
    reimported = manual_measurements(dict(image=str(ims_file)))
    assert len(reimported['cells']) == 1 and len(reimported['points']) == 2


def test_degenerate_imaris_pole_triangles_do_not_bias_volume():
    surface = box_surface()
    surface.vertices = np.vstack([surface.vertices, surface.vertices[[0,0,1]]])
    surface.triangles = np.vstack([surface.triangles, [8,9,10]])
    assert surface_volume(surface) == pytest.approx(24)
