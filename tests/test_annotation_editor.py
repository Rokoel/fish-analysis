"""Regression coverage for live geometry, undo and annotated-image exports."""

from copy import deepcopy
from pathlib import Path

import h5py
import numpy as np
import pytest
from openpyxl import load_workbook
from test_workflow import signal_result  # noqa: F401 - shared calibrated fixture

from fish_viewer.analysis.annotations import AnnotationDocument, origin
from fish_viewer.models import Measurement, Scene, Surface
from fish_viewer.reader import entry_for_path, load_dataset
from fish_viewer.workflow.edit_export import editable_scene_xml, export_document
from fish_viewer.workflow.volumes import surface_volume


def cell():
    vertices = np.array(
        [
            [11, 21, 30],
            [17, 21, 30],
            [17, 25, 30],
            [11, 25, 30],
            [11, 21, 34],
            [17, 21, 34],
            [17, 25, 34],
            [11, 25, 34],
        ],
        float,
    )
    triangles = np.array(
        [
            [0, 2, 1],
            [0, 3, 2],
            [4, 5, 6],
            [4, 6, 7],
            [0, 1, 5],
            [0, 5, 4],
            [1, 2, 6],
            [1, 6, 5],
            [2, 3, 7],
            [2, 7, 6],
            [3, 0, 4],
            [3, 4, 7],
        ]
    )
    return Surface("manual cell", vertices, triangles)


def document(ims_file):
    dataset = load_dataset(entry_for_path(ims_file))
    dataset.scene = Scene(
        surfaces=[cell()],
        measurements=[
            Measurement(
                "green signal; channel 1",
                np.array([[14, 23, 31], [14, 23, 30]], float),
                ["A", "B"],
                [0, 0],
                [(0, 1)],
            )
        ],
    )
    return AnnotationDocument(dataset)


def test_imported_image_and_scene_stay_unchanged_through_edits_and_export(qapp, ims_file, tmp_path):
    image = Path(export_document(document(ims_file), tmp_path / 'imported.ims'))
    originals = {image.with_suffix(suffix): image.with_suffix(suffix).read_bytes()
                 for suffix in ('.ims', '.xlsx')}
    doc = AnnotationDocument(load_dataset(entry_for_path(image)))
    polygon = np.array([[11, 20], [17, 20], [17, 26], [11, 26]], float)
    doc.add_layered_cell({30.: polygon, 34.: polygon}, 2)
    doc.delete('surface', 1)
    doc.undo.undo()
    doc.move('surface', 0, 0, [11, 21, 29])
    export_document(doc, tmp_path / 'edited-copy.ims')
    with pytest.raises(ValueError, match='original scan is read-only'):
        export_document(doc, image)
    with pytest.raises(FileExistsError, match='never overwritten'):
        export_document(doc, tmp_path / 'edited-copy.ims')
    for path, contents in originals.items():
        assert path.read_bytes() == contents


def test_move_signal_contour_delete_and_undo_recalculate(qapp, ims_file):
    doc = document(ims_file)
    assert doc.volumes == [pytest.approx(96)] and doc.distances[0] == pytest.approx(1)
    doc.move("measurement", 0, 0, [14, 23, 31.5])
    assert doc.distances[0] == pytest.approx(1.5)
    assert origin(doc.dataset.scene.measurements[0]) == "Manual"
    assert doc.dataset.scene.measurements[0].edited
    doc.move("surface", 0, 0, [11, 21, 29])
    assert doc.volumes[0] > 96
    assert doc.distances[0] > 1.5
    doc.undo.undo()
    assert doc.volumes[0] == pytest.approx(96)
    doc.undo.undo()
    assert doc.distances[0] == pytest.approx(1)
    assert origin(doc.dataset.scene.measurements[0]) == "Manual"
    doc.delete("surface", 0)
    assert not doc.dataset.scene.surfaces
    assert not doc.dataset.scene.measurements
    assert doc.distances == {}
    doc.undo.undo()
    assert doc.dataset.scene.surfaces and doc.distances[0] == pytest.approx(1)
    doc.delete("measurement", 0)
    assert not doc.dataset.scene.measurements
    doc.undo.undo()
    assert doc.dataset.scene.measurements


def test_add_remove_mesh_point_preserves_closed_volume(qapp, ims_file):
    doc = document(ims_file)
    v = deepcopy(doc.dataset.scene.surfaces[0].vertices)
    doc.add_vertex("surface", 0, [14, 23, 30])
    assert len(doc.dataset.scene.surfaces[0].vertices) == 9
    assert doc.volumes[0] == pytest.approx(96)
    doc.delete_vertex("surface", 0, 8)
    assert len(doc.dataset.scene.surfaces[0].vertices) == 8
    assert doc.volumes[0] == pytest.approx(96)
    doc.undo.undo()
    assert len(doc.dataset.scene.surfaces[0].vertices) == 9
    doc.undo.undo()
    np.testing.assert_array_equal(doc.dataset.scene.surfaces[0].vertices, v)


def test_added_signal_outside_cell_has_no_fabricated_distance(qapp, ims_file):
    doc = document(ims_file)
    doc.add_signal([14, 23, 32], 1)
    assert doc.bindings[1] == 0 and doc.distances[1] == pytest.approx(2)
    doc.move("measurement", 1, 0, [100, 100, 100])
    assert doc.bindings[1] is None and 1 not in doc.distances
    assert doc.dataset.scene.measurements[1].labels == ["A"]
    doc.undo.undo()
    assert doc.distances[1] == pytest.approx(2)


@pytest.mark.parametrize("embedded", [False, True])
def test_export_latest_edited_geometry_and_provenance(
    qapp, ims_file, tmp_path, embedded
):
    doc = document(ims_file)
    if embedded:
        data = editable_scene_xml(
            ims_file.name, doc.dataset.scene, doc.dataset.metadata
        )
        with h5py.File(ims_file, "r+") as f:
            f.create_dataset("Scene/Data", data=np.frombuffer(data, dtype=np.uint8))
        loaded = load_dataset(entry_for_path(ims_file))
        assert len(loaded.scene.surfaces) == len(loaded.scene.measurements) == 1
        assert loaded.scene.surfaces[0].source == "manual"
    source = ims_file.read_bytes()
    doc.add_signal([14, 23, 32], 1)
    doc.move("surface", 0, 0, [11, 21, 29])
    doc.delete("measurement", 0)
    path = Path(export_document(doc, tmp_path / "edited.ims"))
    reopened = load_dataset(entry_for_path(path))
    assert len(reopened.scene.surfaces) == len(reopened.scene.measurements) == 1
    np.testing.assert_allclose(
        reopened.scene.surfaces[0].vertices, doc.dataset.scene.surfaces[0].vertices
    )
    assert reopened.scene.surfaces[0].source == "manual"
    assert reopened.scene.surfaces[0].name == "manual cell"
    assert reopened.scene.measurements[0].source == "manual"
    assert reopened.scene.measurements[0].run_id == doc.run_id
    assert reopened.scene.measurements[0].created_at in reopened.scene.measurements[0].name
    assert ims_file.read_bytes() == source
    with h5py.File(path) as out, h5py.File(ims_file) as original:
        assert 'Scene/Content/Surfaces0' in out and 'Scene/Data' in out
        np.testing.assert_array_equal(
            out["DataSet/ResolutionLevel 0/TimePoint 0/Channel 0/Data"],
            original["DataSet/ResolutionLevel 0/TimePoint 0/Channel 0/Data"],
        )
    workbook = load_workbook(path.with_suffix(".xlsx"))
    assert workbook.sheetnames == ["Measurements"]
    rows = list(workbook.active.values)
    assert (
        len(rows) == 3
        and rows[1][2] == pytest.approx(doc.volumes[0])
        and rows[2][2] == pytest.approx(doc.volumes[0])
        and rows[2][-1] == pytest.approx(doc.distances[0])
    )
    workbook.close()


def test_manual_import_keeps_original_b_until_edited(qapp, ims_file):
    doc = document(ims_file)
    doc.dataset.scene.measurements[0].points[1] = [13.8, 23, 30]
    doc = AnnotationDocument(doc.dataset)
    assert doc.distances[0] == pytest.approx(np.sqrt(1.04))
    doc.move("measurement", 0, 0, [14, 23, 31])
    assert doc.distances[0] == pytest.approx(1)


def test_hdf_binary_spots_meshes_and_vlen_scene_xml(ims_file):
    surface = cell()
    with h5py.File(ims_file, "r+") as f:
        g = f.create_group("Scene8/Content/Surfaces0")
        g.attrs["Name"] = "embedded cell"
        g.create_dataset("Vertices", data=surface.vertices)
        g.create_dataset("Triangles", data=surface.triangles)
        g = f.create_group("Scene8/Content/Points0")
        g.attrs["Name"] = "embedded signals"
        g.create_dataset("CoordsXYZR", data=[[14, 23, 31, 0.1], [14, 23, 32, 0.1]])
        g.create_dataset("Time", data=[[0], [0]])
        xml = f.create_dataset(
            "Scene8/Data", (1,), dtype=h5py.vlen_dtype(np.dtype("uint8"))
        )
        xml[0] = np.frombuffer(
            b"<bpImarisApplication></bpImarisApplication>", dtype=np.uint8
        )
    dataset = load_dataset(entry_for_path(ims_file))
    assert len(dataset.scene.surfaces) == 1 and len(dataset.scene.measurements) == 2
    assert surface_volume(dataset.scene.surfaces[0]) == pytest.approx(96)
    assert not dataset.scene.warnings


@pytest.mark.gui
def test_editor_ui_manual_import_selection_edit_undo_and_hide(qapp, ims_file, tmp_path):
    from PySide6 import QtCore

    from fish_viewer.analysis.app import AnalysisWindow

    doc = document(ims_file)
    window = AnalysisWindow(
        off_screen=True,
        settings=QtCore.QSettings(
            str(tmp_path / "settings.ini"), QtCore.QSettings.Format.IniFormat
        ),
    )
    try:
        window.set_dataset(doc.dataset, None)
        assert window.pipeline.currentData() == "consensus"
        assert window.pipeline.count() == 2
        editor = window.editor
        assert editor.tree.topLevelItem(0).childCount() == 2
        item = editor.tree.topLevelItem(0).child(0)
        editor.tree.setCurrentItem(item)
        editor.move_point(0, [11, 21, 29])
        assert editor.document.volumes[0] > 96
        editor.document.undo.undo()
        assert editor.document.volumes[0] == pytest.approx(96)
        editor.manual.setChecked(False)
        assert not window.dataset.scene.surfaces[0].visible
        editor.delete()
        assert not window.dataset.scene.surfaces
        editor.document.undo.undo()
        assert window.dataset.scene.surfaces
        editor.edit.setChecked(True)
        window.change_slice(0)
        qapp.processEvents()
    finally:
        window.close()


def test_slice_border_point_move_is_closed_and_undoable(qapp, ims_file):
    doc = document(ims_file)
    # The point lies on an edge shared by two triangles: both must be split.
    doc.move_section(0, [14, 21, 30], [14, 20.8, 30])
    assert doc.volumes[0] is not None
    assert len(doc.dataset.scene.surfaces[0].vertices) == 9
    doc.undo.undo()
    assert doc.volumes[0] == pytest.approx(96)
    assert len(doc.dataset.scene.surfaces[0].vertices) == 8


@pytest.mark.gui
def test_contour_edit_updates_volumes_distances_and_tables_with_undo(qapp, ims_file, tmp_path):
    from PySide6 import QtCore
    from fish_viewer.analysis.app import AnalysisWindow

    doc = document(ims_file)
    other = cell()
    other.vertices += [20, 0, 0]
    doc.dataset.scene.surfaces.append(other)
    doc.dataset.scene.measurements.append(Measurement(
        'Unedited measurement', np.array([[34, 23, 31], [34, 23, 30.2]], float),
        ['A', 'B'], [0, 0], [(0, 1)]))
    window = AnalysisWindow(off_screen=True, settings=QtCore.QSettings(
        str(tmp_path / 'measurements.ini'), QtCore.QSettings.Format.IniFormat))
    try:
        window.set_dataset(doc.dataset)
        live = window.editor.document
        before = live.dataset.scene.measurements[1].points.copy()
        live.move_section(0, [14, 23, 30], [14, 23, 29], axis=2)
        # The displaced bottom face adds a pyramid of volume 6*4*1/3.
        assert live.volumes == pytest.approx([104, 96])
        assert live.distances == pytest.approx({0: np.sqrt(3.2), 1: .8})
        assert float(window.counts.item(0, 4).text()) == pytest.approx(104)
        assert float(window.distances.item(0, 3).text()) == pytest.approx(np.sqrt(3.2), abs=5e-5)
        assert live.report()['manual_measurements']['cells'][0]['volume_um3'] == pytest.approx(104)
        np.testing.assert_array_equal(live.dataset.scene.measurements[1].points, before)
        live.undo.undo()
        assert live.volumes == pytest.approx([96, 96])
        assert live.distances == pytest.approx({0: 1, 1: .8})
        assert float(window.counts.item(0, 4).text()) == pytest.approx(96)
        assert float(window.distances.item(0, 3).text()) == pytest.approx(1)
        live.undo.redo()
        assert live.volumes[0] == pytest.approx(104)
        assert live.distances[0] == pytest.approx(np.sqrt(3.2))
    finally:
        window.close()


@pytest.mark.parametrize("axis,coordinate", [(0, 14), (1, 23), (2, 32)])
def test_repeated_slice_edits_keep_neighbor_points_fixed(
    qapp, ims_file, axis, coordinate
):
    from fish_viewer.views import mesh_for, slice_segments

    doc = document(ims_file)
    surface = doc.dataset.scene.surfaces[0]
    section = np.unique(
        slice_segments(mesh_for(surface), axis, coordinate).reshape(-1, 3), axis=0
    )
    old = section[0].copy()
    doc.move_section(0, old, old, axis=axis)
    assert doc.volumes[0] == pytest.approx(96)
    assert mesh_for(doc.dataset.scene.surfaces[0]).n_open_edges == 0
    vertex_count = len(doc.dataset.scene.surfaces[0].vertices)
    for _ in range(2):
        new = old.copy()
        new[(axis + 1) % 3] -= 0.2
        doc.move_section(0, old, new, axis=axis)
        surface = doc.dataset.scene.surfaces[0]
        actual = np.unique(
            slice_segments(mesh_for(surface), axis, coordinate).reshape(-1, 3), axis=0
        )
        section[np.linalg.norm(section - old, axis=1) < 1e-7] = new
        assert len(actual) == len(section)
        for point in section:
            assert np.linalg.norm(actual - point, axis=1).min() < 1e-7
        assert len(surface.vertices) == vertex_count
        assert mesh_for(surface).n_open_edges == 0
        old = new
    doc.undo.undo()
    doc.undo.undo()
    assert doc.volumes[0] == pytest.approx(96)
    doc.undo.undo()
    assert len(doc.dataset.scene.surfaces[0].vertices) == 8


@pytest.mark.gui
@pytest.mark.parametrize("axis,index", [(0, 0), (1, 1), (2, 1)])
@pytest.mark.parametrize(
    "zoom,click_offset", [(0.5, 0), (2.0, 0), (0.5, 9), (2.0, 9), (0.1, 0)]
)
def test_drag_slice_contour_with_mouse(
    qapp, ims_file, tmp_path, axis, index, zoom, click_offset
):
    from PySide6 import QtCore, QtGui, QtTest

    from fish_viewer.analysis.app import AnalysisWindow

    window = AnalysisWindow(
        off_screen=True,
        settings=QtCore.QSettings(
            str(tmp_path / "drag.ini"), QtCore.QSettings.Format.IniFormat
        ),
    )
    try:
        window.resize(1300, 900)
        window.show()
        window.set_dataset(document(ims_file).dataset, None)
        editor, view = window.editor, window.slice_view
        editor.tree.setCurrentItem(editor.tree.topLevelItem(0).child(0))
        assert editor.edit.isChecked()
        editor.set_tool_mode("precise")
        window.change_axis(axis)
        window.change_slice(index)
        view.fit()
        view.scale(zoom, zoom)
        qapp.processEvents()
        assert editor.handles
        np.testing.assert_allclose(
            [h.xyz[2 - axis] for h in editor.handles], view.coordinate
        )
        positions = np.array([h.xyz for h in editor.handles])
        assert len(np.unique(positions, axis=0)) == len(positions)
        handle = next(h for h in editor.handles if h.vertex is None)
        original = handle.xyz.copy()
        before = window.dataset.scene.surfaces[0].vertices.copy()
        start = view.mapFromScene(handle.pos()) + QtCore.QPoint(click_offset, 0)
        end = start + QtCore.QPoint(-20, -10)
        expected = original.copy()
        delta = view.mapToScene(end) - view.mapToScene(start)
        u, v = view.projection_axes
        expected[u] += delta.x()
        expected[v] += delta.y()
        assert handle in view.items(start)
        QtTest.QTest.mousePress(
            view.viewport(), QtCore.Qt.MouseButton.LeftButton, pos=start
        )
        assert view.scene().mouseGrabberItem() is handle
        move = QtGui.QMouseEvent(
            QtCore.QEvent.Type.MouseMove,
            QtCore.QPointF(end),
            QtCore.QPointF(view.viewport().mapToGlobal(end)),
            QtCore.Qt.MouseButton.NoButton,
            QtCore.Qt.MouseButton.LeftButton,
            QtCore.Qt.KeyboardModifier.NoModifier,
        )
        qapp.sendEvent(view.viewport(), move)
        # Lines must follow during the drag, before changing saved geometry.
        np.testing.assert_array_equal(window.dataset.scene.surfaces[0].vertices, before)
        overlay = next(
            item for item in view._overlays if item.data(0) == ("surface", 0)
        )
        preview = overlay.path()
        assert any(
            np.allclose(
                [preview.elementAt(i).x, preview.elementAt(i).y], expected[[u, v]]
            )
            for i in range(preview.elementCount())
        )
        QtTest.QTest.mouseRelease(
            view.viewport(), QtCore.Qt.MouseButton.LeftButton, pos=end
        )
        qapp.processEvents()
        after = window.dataset.scene.surfaces[0].vertices
        assert not np.array_equal(before, after)
        # Introducing slice vertices must preserve all original mesh vertices
        # and every other displayed border point.
        np.testing.assert_array_equal(after[: len(before)], before)
        assert np.linalg.norm(after - expected, axis=1).min() < 1e-7
        remaining = positions[np.linalg.norm(positions - original, axis=1) > 1e-6]
        displayed = np.array([h.xyz for h in editor.handles])
        assert len(displayed) == len(positions)
        for point in remaining:
            assert np.linalg.norm(displayed - point, axis=1).min() < 1e-6
        from fish_viewer.views import mesh_for

        assert mesh_for(window.dataset.scene.surfaces[0]).n_open_edges == 0
        assert editor.document.undo.count() == 1
        assert editor.document.volumes[0] != pytest.approx(96)
        assert len(window.dataset.scene.measurements) == 1
        assert editor.document.dataset.scene.surfaces[0].source == "manual"
        assert editor.document.dataset.scene.surfaces[0].edited
        view.setFocus()
        QtTest.QTest.keyClick(
            view.viewport(),
            QtCore.Qt.Key.Key_Z,
            QtCore.Qt.KeyboardModifier.ControlModifier,
        )
        qapp.processEvents()
        np.testing.assert_array_equal(window.dataset.scene.surfaces[0].vertices, before)
    finally:
        window.close()


def test_multiple_spot_centers_are_individually_deletable(qapp, ims_file):
    doc = document(ims_file)
    doc.dataset.scene.measurements = [
        Measurement(
            "Spots", np.array([[14, 23, 31], [14, 23, 32]], float), ["A", "A"], [0, 0]
        )
    ]
    doc = AnnotationDocument(doc.dataset)
    assert len(doc.dataset.scene.measurements) == 2
    doc.delete("measurement", 0)
    assert len(doc.dataset.scene.measurements) == 1
    assert doc.dataset.scene.measurements[0].points[0, 2] == 32
    doc.undo.undo()
    assert len(doc.dataset.scene.measurements) == 2


@pytest.mark.gui
def test_edits_and_undo_survive_switching_scans(qapp, ims_file, tmp_path):
    from PySide6 import QtCore

    from fish_viewer.analysis.app import AnalysisWindow

    doc = document(ims_file)
    window = AnalysisWindow(
        off_screen=True,
        settings=QtCore.QSettings(
            str(tmp_path / "settings.ini"), QtCore.QSettings.Format.IniFormat
        ),
    )
    try:
        window.set_dataset(doc.dataset, None)
        window.editor.document.delete("surface", 0)
        other = deepcopy(doc.dataset)
        other.metadata.path = tmp_path / "other.ims"
        window.set_dataset(other, None)
        assert window.documents[window.document_key(doc.dataset)].dataset.volumes == []
        window.set_dataset(load_dataset(entry_for_path(ims_file)), None)
        assert not window.dataset.scene.surfaces
        window.editor.document.undo.undo()
        assert len(window.dataset.scene.surfaces) == 1
    finally:
        window.close()


def test_pipeline_export_preserves_embedded_manual_and_new_detections(
    qapp,
    signal_result,  # noqa: F811 - imported pytest fixture
    tmp_path,
):
    from fish_viewer.workflow.distances import measure_result
    from fish_viewer.workflow.export import export_result
    from fish_viewer.workflow.results import filter_result

    filtered = filter_result(signal_result, tmp_path / "work")
    measured = measure_result(
        Path(filtered["directory"]) / "report.json", tmp_path / "work"
    )
    image = Path(measured["image"])
    original = load_dataset(entry_for_path(image))
    manual = Scene(
        surfaces=[cell()],
        measurements=[
            Measurement(
                "original A-B",
                np.array([[14, 23, 31], [14, 23, 30]], float),
                ["A", "B"],
                [0, 0],
                [(0, 1)],
            )
        ],
    )
    with h5py.File(image, "r+") as file:
        file.create_dataset(
            "Scene/Data",
            data=np.frombuffer(
                editable_scene_xml(image.name, manual, original.metadata),
                dtype=np.uint8,
            ),
        )
    target = tmp_path / "combined.ims"
    exported = export_result(Path(measured["directory"]) / "report.json", target)
    reopened = load_dataset(entry_for_path(target))
    assert all(s.source == "manual" for s in reopened.scene.surfaces)
    from fish_viewer.annotation_state import activate_run
    activate_run(reopened.scene, exported["annotation_run_id"])
    assert sum(s.source == "manual" for s in reopened.scene.surfaces) == 1
    assert sum(s.source != "manual" for s in reopened.scene.surfaces) == 2
    assert sum(m.source == "manual" for m in reopened.scene.measurements) == 1
    assert sum(m.source != "manual" for m in reopened.scene.measurements) == 2
    assert not reopened.scene.warnings


def test_deleting_cell_does_not_change_unassociated_manual_ab(qapp, ims_file):
    doc = document(ims_file)
    orphan = Measurement(
        "outside manually measured pair",
        np.array([[100, 100, 100], [100, 100, 101]], float),
        ["A", "B"],
        [0, 0],
        [(0, 1)],
    )
    doc.dataset.scene.measurements.append(orphan)
    doc = AnnotationDocument(doc.dataset)
    assert doc.bindings[1] is None and doc.distances[1] == pytest.approx(1)
    doc.delete("surface", 0)
    assert len(doc.dataset.scene.measurements) == 1
    assert doc.dataset.scene.measurements[0].uid == orphan.uid
    assert doc.dataset.scene.measurements[0].labels == ["A", "B"]
    assert doc.distances[0] == pytest.approx(1)


def test_unsupported_annotations_cannot_be_silently_discarded(qapp, ims_file, tmp_path):
    doc = document(ims_file)
    doc.dataset.scene.warnings = [
        "bpSpots annotations are not supported by this XML reader."
    ]
    with pytest.raises(ValueError, match="cannot safely discard"):
        export_document(doc, tmp_path / "unsafe.ims")
    assert not (tmp_path / "unsafe.ims").exists()


def test_reopening_preserves_edited_signal_association_and_cleared_score(
    qapp, ims_file, tmp_path
):
    doc = document(ims_file)
    doc.dataset.scene.surfaces[0].name = "Predicted cell 1 - 1 signals"
    second = deepcopy(doc.dataset.scene.surfaces[0])
    from uuid import uuid4
    second.uid = str(uuid4())
    second.name = "Predicted cell 2 - 0 signals"
    second.vertices += [10, 0, 0]
    doc.dataset.scene.surfaces.append(second)
    doc.dataset.scene.measurements[
        0
    ].name = "Signal 1; cell 1; channel 2; score .99; distance 1 um"
    doc.channels[0] = 1
    doc.scores[0] = 0.99
    doc.move("measurement", 0, 0, [24, 23, 31])
    assert doc.bindings[0] == 1 and doc.scores[0] is None
    assert doc.dataset.scene.measurements[0].cell_uid == second.uid
    assert doc.dataset.scene.measurements[0].score is None
    exported = export_document(doc, tmp_path / "moved.ims")
    reopened = AnnotationDocument(load_dataset(entry_for_path(exported)))
    assert reopened.bindings[0] == 1 and reopened.scores[0] is None
    assert reopened.distances[0] == pytest.approx(doc.distances[0])
    reopened.move("measurement", 0, 0, [100, 100, 100])
    exported = export_document(reopened, tmp_path / "outside.ims")
    again = AnnotationDocument(load_dataset(entry_for_path(exported)))
    assert again.bindings[0] is None and not again.distances
    assert again.dataset.scene.measurements[0].labels == ["A"]


def test_embedded_annotation_metadata_is_ignored_if_scene_changed(qapp, ims_file, tmp_path):
    import json

    doc = document(ims_file)
    exported = Path(export_document(doc, tmp_path / "saved.ims"))
    with h5py.File(exported, 'r+') as image:
        data = json.loads(image['FISHAnalysis/Annotations'][()].tobytes())
        data['editor_state']['objects']['measurements'][0]['cell_uid'] = None
        del image['FISHAnalysis/Annotations']
        image.create_dataset('FISHAnalysis/Annotations', data=np.frombuffer(json.dumps(data).encode(), np.uint8))
        image['Scene/Data'][0] = np.concatenate((image['Scene/Data'][0], np.frombuffer(b'\n', np.uint8)))
    reopened = AnnotationDocument(load_dataset(entry_for_path(exported)))
    assert reopened.bindings[0] == 0
