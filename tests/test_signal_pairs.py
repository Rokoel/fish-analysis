from copy import deepcopy
from pathlib import Path

import h5py
import numpy as np
import pytest
from openpyxl import load_workbook

from fish_viewer.analysis.annotations import AnnotationDocument
from fish_viewer.models import Channel, Measurement
from fish_viewer.reader import entry_for_path, load_dataset
from fish_viewer.workflow.edit_export import export_document
from fish_viewer.workflow.signal_pairs import pair_signals
from test_annotation_editor import document


def point(name, channel, xyz, cell='cell', time=0):
    return dict(signal_id=name, channel=channel, xyz=xyz, cell_id=cell, time=time)


def test_closest_first_pairing_is_unique_per_channel_pair_with_physical_cutoff():
    points = [point('a', 0, [0, 0, 0]), point('b', 0, [.3, 0, 0]),
              point('c', 1, [.1, 0, 0]), point('d', 1, [.8, 0, 0]),
              point('e', 2, [.05, 0, 0]), point('other-cell', 1, [0, 0, 0], cell='other'),
              point('other-time', 1, [0, 0, 0], time=1), point('outside', 1, [0, 0, 0], cell=None)]
    pairs = pair_signals(points, .5)
    channel01 = [p for p in pairs if (p['channel_a'], p['channel_b']) == (0, 1)]
    assert [(p['signal_id_a'], p['signal_id_b']) for p in channel01] == [('a', 'c'), ('b', 'd')]
    assert channel01[1]['distance_um'] == pytest.approx(.5)
    assert all(p['cell_id'] == 'cell' and p['time'] == 0 and p['distance_um'] <= .5 for p in pairs)
    assert any(p['signal_id_a'] == 'a' and p['signal_id_b'] == 'e' for p in pairs)
    assert len({p['pair_id'] for p in pairs}) == len(pairs)
    assert pair_signals(list(reversed(points)), .5) == pairs
    assert pair_signals(points[:1], .5) == []
    for invalid in (0, -1, float('inf'), float('nan')):
        with pytest.raises(ValueError):
            pair_signals(points, invalid)


def paired_document(ims_file):
    with h5py.File(ims_file, 'r+') as image:
        for level in (0, 1):
            for time in (0, 1):
                group = image[f'DataSet/ResolutionLevel {level}/TimePoint {time}']
                image.copy(group['Channel 1'], group, name='Channel 2')
        image.copy('DataSetInfo/Channel 1', 'DataSetInfo/Channel 2')
        image['DataSetInfo/Channel 2'].attrs['Name'] = np.frombuffer(b'Red', dtype='S1')
    doc = document(ims_file)
    doc.channels[0] = 1
    doc.recalculate()
    doc.add_signal([14.4, 23, 31], 2)
    return doc


def test_pairs_update_with_moves_cutoff_cell_edits_deletion_and_undo(qapp, ims_file):
    doc = paired_document(ims_file)
    assert len(doc.pairs) == 1
    original = deepcopy(doc.pairs[0])
    assert original['distance_um'] == pytest.approx(.4)
    for action in (lambda: doc.move('measurement', 2, 0, [15, 23, 31]),
                   lambda: doc.translate({('measurement', 2)}, [1, 0, 0]),
                   lambda: doc.delete('measurement', 2)):
        with pytest.raises(ValueError, match='calculated'):
            action()
    assert doc.pairs[0] == original
    doc.move('measurement', 1, 0, [15.5, 23, 31])
    assert doc.pairs == []
    doc.undo.undo()
    assert doc.pairs[0]['pair_id'] == original['pair_id']
    doc.set_pairing_distance(.3)
    assert doc.pairs == []
    doc.undo.undo()
    assert len(doc.pairs) == 1 and doc.pairing_max_distance == 1
    doc.delete('measurement', 1)
    assert not doc.pairs
    doc.undo.undo()
    assert len(doc.pairs) == 1
    def shrink_cell():
        surface = doc.dataset.scene.surfaces[0]
        surface.vertices[:, 0] = 14 + (surface.vertices[:, 0] - 14) * .05
        surface.edited = True
    doc.apply('Shrink contour', shrink_cell)
    assert not doc.pairs
    doc.undo.undo()
    assert len(doc.pairs) == 1
    doc.delete('surface', 0)
    assert not doc.pairs and not doc.dataset.scene.measurements
    doc.undo.undo()
    assert len(doc.pairs) == 1


def test_standalone_ims_and_excel_include_pairs_and_preserve_pixels(qapp, ims_file, tmp_path):
    doc = paired_document(ims_file)
    doc.dataset.scene.surfaces[0].opacity = .71
    before = ims_file.read_bytes()
    output = tmp_path / 'output'
    output.mkdir()
    path = Path(export_document(doc, output / 'paired.ims'))
    assert {p.name for p in output.iterdir()} == {'paired.ims', 'paired.xlsx'}
    assert ims_file.read_bytes() == before
    loaded = load_dataset(entry_for_path(path))
    assert not loaded.scene.warnings
    assert len(loaded.scene.surfaces) == 1
    assert loaded.scene.surfaces[0].opacity == pytest.approx(.71)
    lines = [m for m in loaded.scene.measurements if m.kind == 'signal_pair']
    assert len(lines) == 1 and lines[0].length == pytest.approx(.4)
    reopened = AnnotationDocument(loaded)
    assert len(reopened.pairs) == 1 and reopened.pairs[0]['pair_id'] == doc.pairs[0]['pair_id']
    assert len(reopened.dataset.scene.measurements) == 3
    workbook = load_workbook(path.with_suffix('.xlsx'))
    assert workbook.sheetnames == ['Measurements', 'Signal pairs']
    row = list(workbook['Signal pairs'].values)[1]
    assert row[4] == 2 and row[7] == 3 and row[-2] == pytest.approx(.4) and row[-1] == 1
    workbook.close()
    with h5py.File(ims_file) as source, h5py.File(path) as result:
        import xml.etree.ElementTree as ET
        root = ET.fromstring(result['Scene/Data'][0].tobytes())
        viewer = next(root.iter('bpSurfacesViewer'))
        assert viewer.find('bpSurfacesId').get('Value') == str(result['Scene/Content/Surfaces0'].attrs['Id'])
        assert viewer.find('bpSurfaces/bpSurfacesId') is None
        for level in (0, 1):
            for time in (0, 1):
                for channel in (0, 1, 2):
                    key = f'DataSet/ResolutionLevel {level}/TimePoint {time}/Channel {channel}/Data'
                    np.testing.assert_array_equal(source[key], result[key])
        assert result['Scene/Content/Surfaces0/Vertices'].shape[1] == 6
        assert result['Scene/Data'].dtype.kind == 'O'


def test_pair_cutoff_and_roundness_controls(qapp, ims_file, tmp_path):
    from PySide6 import QtCore
    from fish_viewer.analysis.app import AnalysisWindow
    settings = QtCore.QSettings(str(tmp_path / 'preferences.ini'), QtCore.QSettings.Format.IniFormat)
    settings.setValue('minimum_roundness_percent', 72)
    window = AnalysisWindow(settings=settings)
    try:
        assert window.roundness.value() == 82 and window.roundness.minimum() == 75
        window.roundness.setValue(50)
        assert window.roundness.value() == 75
        doc = paired_document(ims_file)
        window.set_dataset(doc.dataset, None)
        assert len(window.editor.document.pairs) == 1
        assert window.distances.rowCount() == 3
        window.pair_distance.setValue(.3)
        assert not window.editor.document.pairs
        assert window.distances.rowCount() == 2
        window.editor.document.undo.undo()
        assert window.pair_distance.value() == 1 and window.distances.rowCount() == 3
    finally:
        window.close()


@pytest.mark.integration
def test_legacy_imaris_image_exports_as_standalone_hdf5_without_voxel_changes(qapp, reference, tmp_path):
    source = reference.parents[1] / 'FISH antiLAD/18327antiLAD2.ims'
    if not source.exists():
        pytest.skip('Local legacy Imaris example is unavailable')
    before = source.read_bytes()
    original = load_dataset(entry_for_path(source))
    doc = AnnotationDocument(original)
    target = Path(export_document(doc, tmp_path / 'legacy-converted.ims'))
    assert h5py.is_hdf5(target)
    loaded = load_dataset(entry_for_path(target))
    np.testing.assert_allclose(loaded.metadata.spacing, original.metadata.spacing)
    np.testing.assert_allclose(loaded.metadata.origin, original.metadata.origin)
    for a, b in zip(original.volumes, loaded.volumes):
        np.testing.assert_array_equal(a, b)
    assert len(loaded.scene.surfaces) == len(original.scene.surfaces)
    assert len(loaded.scene.measurements) == len(original.scene.measurements)
    assert source.read_bytes() == before
