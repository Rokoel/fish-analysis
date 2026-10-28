"""Installed-binary smoke test, deliberately independent of the checkout."""

import argparse
import json
import os
import time
from dataclasses import asdict
from pathlib import Path


def main(arguments=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument('--device', choices=['cpu', 'gpu', 'auto'], default='cpu')
    parser.add_argument('--scan', type=Path, help='Also run a real scan through the GUI queue and validate its exports')
    args = parser.parse_args(arguments)
    output = args.output.resolve()
    output.mkdir(parents=True, exist_ok=True)
    os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
    os.environ["QT_API"] = "pyside6"
    import h5py
    import numpy as np
    from openpyxl import load_workbook
    from PySide6 import QtCore, QtWidgets

    from .analysis.app import AnalysisWindow
    from .analysis.controller import (
        default_models,
        default_proposal_model,
        default_refinement_model,
    )
    from .neural.config import NeuralConfig
    from .reader import entry_for_path, load_dataset
    from .signals.engine import image_fingerprint
    from .workflow.distances import measure_result
    from .workflow.export import export_result
    from .workflow.results import filter_result

    image = output / "synthetic.ims"
    shape = (64, 64, 64)
    with h5py.File(image, "w") as f:
        info = f.create_group("DataSetInfo/Image")
        for key, value in {
            "ExtMin0": "0",
            "ExtMin1": "0",
            "ExtMin2": "0",
            "ExtMax0": "9.6",
            "ExtMax1": "9.6",
            "ExtMax2": "9.6",
            "Unit": "um",
        }.items():
            info.attrs[key] = np.frombuffer(value.encode(), dtype="S1")
        for c, color in enumerate(("1 0 1", "0 1 0")):
            channel = f.create_group(f"DataSetInfo/Channel {c}")
            for key, value in {
                "Name": ("Cell membrane", "Signal")[c],
                "Color": color,
                "ColorRange": "0 1000",
                "GammaCorrection": "1",
            }.items():
                channel.attrs[key] = np.frombuffer(value.encode(), dtype="S1")
            group = f.create_group(f"DataSet/ResolutionLevel 0/TimePoint 0/Channel {c}")
            for axis, n in zip("ZYX", shape):
                group.attrs[f"ImageSize{axis}"] = str(n).encode()
            group.create_dataset(
                "Data", data=np.zeros(shape, np.uint16), compression="gzip"
            )
    app = QtWidgets.QApplication.instance() or QtWidgets.QApplication([])
    settings = QtCore.QSettings(
        str(output / "settings.ini"), QtCore.QSettings.Format.IniFormat
    )
    settings.setValue('device', 'gpu')
    window = AnalysisWindow(off_screen=True, settings=settings)
    try:
        assert window.pipeline.currentData() == "consensus"
        from .runtime_profile import packaged_profile
        profile = packaged_profile()
        if profile is not None:
            assert window.device.currentData() == ('gpu' if profile['cuda'] else 'cpu')
        cells, signals = default_models()
        done, errors = [], []
        # The controller is driven directly here; its GUI queue callbacks expect
        # a user-selected scan, so exercise that UI separately below.
        window.controller.completed.disconnect()
        window.controller.progress.disconnect()
        window.controller.failed.disconnect()
        window.controller.completed.connect(done.append)
        window.controller.failed.connect(errors.append)
        # Test the real QProcess protocol through every default stage, including
        # the TensorFlow -> standalone PyTorch dispatch, without any FISH scans.
        window.controller.start(
            image,
            output / "pipeline",
            cells,
            signals,
            device=args.device,
            refinement=default_refinement_model(),
            stardist_proposals=True,
            consensus_model=default_proposal_model(),
            roundness=0.82,
            cell_channel=0,
            signal_channels=[1],
        )
        deadline = time.monotonic() + 600
        while not done and not errors and time.monotonic() < deadline:
            app.processEvents()
            time.sleep(0.02)
        if errors:
            raise RuntimeError(errors[0])
        assert done, "Frozen consensus pipeline timed out."
        assert not window.controller.busy
        # Non-empty synthetic contour + A: verify bundled Imaris templates,
        # physical distances, geometry round-trip and the single-sheet workbook.
        dataset = load_dataset(entry_for_path(image))
        labels = np.zeros(shape, np.uint32)
        labels[16:48, 16:48, 16:48] = 1
        metadata = {
            "version": 1,
            "image": str(image),
            "image_name": image.name,
            "time": 0,
            "origin_xyz": dataset.center_origin,
            "spacing_zyx": dataset.spacing[::-1],
            "channel": 0,
            "channel_scores": {},
            "fingerprint": image_fingerprint(dataset),
            "config": asdict(NeuralConfig()),
            "cells": [{"label": 1}],
        }
        prediction = output / "contours.npz"
        np.savez_compressed(
            prediction, labels=labels, metadata=np.array(json.dumps(metadata))
        )
        report = {
            "version": 1,
            "engine": "signal-center-unet3d",
            "mode": "apply",
            "status": "complete",
            "directory": str(output),
            "image": str(image),
            "scene": None,
            "time": 0,
            "cell_predictions": str(prediction),
            "fingerprint": metadata["fingerprint"],
            "channels": [
                {"name": c.name, "color": c.color} for c in dataset.metadata.channels
            ],
            "channel_index_base": 0,
            "unit": "µm",
            "membrane_channel": 0,
            "counts": [],
            "points": [
                {
                    "signal_id": 1,
                    "cell_id": 1,
                    "channel": 1,
                    "xyz": [4.8, 4.8, 4.8],
                    "score": 0.99,
                }
            ],
        }
        report_path = output / "report.json"
        report_path.write_text(json.dumps(report))
        filtered = filter_result(report_path, output / "nonempty")
        measured = measure_result(
            Path(filtered["directory"]) / "report.json", output / "nonempty"
        )
        target = output / "measured.ims"
        exported = export_result(Path(measured["directory"]) / "report.json", target)
        reopened = load_dataset(entry_for_path(target))
        assert (
            len(reopened.scene.surfaces) == 1 and len(reopened.scene.measurements) == 1
        )
        assert measured["cell_volumes"][0]["volume_um3"] > 0
        assert measured["points"][0]["distance_um"] > 0
        books = list(output.rglob("*.xlsx"))
        assert books, exported
        with_books = load_workbook(books[-1], read_only=True)
        assert with_books.sheetnames == ["Measurements"]
        with_books.close()
        window.slice_view.set_dataset(reopened)
        window.slice_view.update_image()
        print(
            json.dumps(
                {
                    "status": "complete",
                    "pipeline": "consensus",
                    "stages": list(window.controller.state["stages"]),
                    "volume_um3": measured["cell_volumes"][0]["volume_um3"],
                    "distance_um": measured["points"][0]["distance_um"],
                    "export": exported,
                }
            )
        )
        if args.scan:
            window.controller.completed.disconnect()
            window.controller.failed.disconnect()
            window.controller.completed.connect(window.completed)
            window.controller.failed.connect(window.failed)
            window.controller.progress.connect(window.processing_progress)
            window.error = errors.append
            window.output.setText(str(output / 'real-scan'))
            window.device.setCurrentIndex(window.device.findData(args.device))
            window.show()
            window.add_paths([args.scan.resolve()])
            deadline = time.monotonic() + 1800
            while (window.dataset is None or window.loading or window._jobs) and not errors and time.monotonic() < deadline:
                app.processEvents()
                time.sleep(.02)
            assert window.dataset is not None and not window.loading and not errors, errors
            window.start()
            while (window.running or window.loading or window._jobs) and time.monotonic() < deadline:
                app.processEvents()
                time.sleep(.02)
            assert not window.running and not window.loading and not errors, errors
            record = window.scans.item(0).data(QtCore.Qt.ItemDataRole.UserRole)
            assert record['status'] == 'completed', record.get('error')
            result = json.loads(Path(record['result']).read_text())
            assert result['status'] == 'complete'
            assert window.result_report['points'], 'The real scan produced no signals.'
            assert window.result_report['cell_volumes'], 'The real scan produced no retained cells.'
            exports = Path(record['result']).parent / 'export'
            workbooks = list(exports.glob('*.xlsx'))
            assert workbooks and list(exports.glob('*.ims'))
            workbook = load_workbook(workbooks[0], read_only=True)
            assert workbook['Measurements'].max_row > 1
            workbook.close()
            window.tabs.setCurrentIndex(3)
            app.processEvents()
            assert window.grab().save(str(output / 'packaged-results.png'))
            window.tabs.setCurrentIndex(0)
            window.slice_view.fit()
            app.processEvents()
            assert window.grab().save(str(output / 'packaged-analysis.png'))
            print(json.dumps(dict(status='complete', real_scan=str(args.scan), cells=len(window.result_report['cell_volumes']), signals=len(window.result_report['points']), report=record['result'])))
    finally:
        window.close()
        app.processEvents()
    return 0
